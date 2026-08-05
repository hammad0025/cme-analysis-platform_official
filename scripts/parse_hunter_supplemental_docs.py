#!/usr/bin/env python3
"""Parse Dr. Hunter's Word/PowerPoint documents that the PDF runs never touched.

The historical parsing runs covered PDFs (cme_complete_parsing_log.json) and
the 25 lettered Reference A-Y methodology docs
(hunter_methodology_knowledge_base.json). Everything else Hunter shared as
.docx/.pptx -- notably the cross-examination playbook docs ("common defense
opinions", "defense doctor testing", spine/TBI memos) -- was skipped.

This script scans the content folders for .docx/.pptx, skips the 25 already
in the methodology KB, dedupes identical copies by content hash (Hunter's
folders mirror the same doc in several places), and writes
hunter_supplemental_knowledge_base.json using the same document shape as
hunter_methodology_knowledge_base.json.
"""
import hashlib
import json
import sys
from pathlib import Path

import docx
import pptx

REPO = Path(__file__).resolve().parents[1]
CONTENT_DIRS = [
    "CME Video Recordings",
    "ortho exam Oregon Hunter",
    "Spine Imaging",
    "TBI ITON articles",
    "TBI Imaging MRI DTI WMH",
    "TBI vestibular",
    "disc & facet trauma",
    "tbi endocrine",
    "tbi prognosis",
]
OUT = REPO / "hunter_supplemental_knowledge_base.json"
MIN_CHARS = 100


def read_docx(path: Path) -> str:
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(p for p in parts if p.strip())


def read_pptx(path: Path) -> str:
    pres = pptx.Presentation(str(path))
    parts = []
    for slide in pres.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(p for p in parts if p.strip())


def main() -> int:
    already = set()
    kb_path = REPO / "hunter_methodology_knowledge_base.json"
    if kb_path.exists():
        kb = json.loads(kb_path.read_text())
        already = {d["filename"].strip().lower() for d in kb.get("documents", [])}

    candidates = []
    for top in CONTENT_DIRS:
        base = REPO / top
        if base.exists():
            candidates += sorted(base.rglob("*.docx")) + sorted(base.rglob("*.pptx"))

    seen_hashes = {}
    documents = []
    skipped_dupes = failed = 0
    for path in candidates:
        if path.name.startswith("~$") or path.name.strip().lower() in already:
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest in seen_hashes:
            skipped_dupes += 1
            continue
        seen_hashes[digest] = path
        try:
            text = read_docx(path) if path.suffix.lower() == ".docx" else read_pptx(path)
        except Exception as e:
            failed += 1
            print(f"FAIL {path.relative_to(REPO)}: {e}")
            continue
        text = text.strip()
        if len(text) < MIN_CHARS:
            failed += 1
            print(f"EMPTY {path.relative_to(REPO)}")
            continue
        documents.append(
            {
                "filename": path.name,
                "path": str(path.relative_to(REPO)),
                "topic": path.parent.name,
                "content_length": len(text),
                "content": text,
                "sha256": digest,
            }
        )
        print(f"OK   {len(text):>7} {path.relative_to(REPO)}")

    OUT.write_text(
        json.dumps(
            {
                "source": "Dr. Hunter supplemental documents (docx/pptx)",
                "description": "Cross-exam playbook, spine/TBI memos, and other "
                "non-PDF documents missed by the original parsing runs",
                "documents": documents,
            },
            indent=1,
        )
    )
    print(
        f"\nDONE: {len(documents)} unique docs parsed, {skipped_dupes} duplicate "
        f"copies skipped, {failed} failed/empty -> {OUT.name}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
