#!/usr/bin/env python3
"""
Parse ALL of Dr. Hunter's methodology documents (A-Y) into the knowledge base.
This is the gold standard for CME analysis.
"""

import os
import json
from docx import Document
from pathlib import Path

CME_DIR = "/Users/hammadhaque/Documents/cme-analysis-platform/CME Video Recordings"

def parse_docx(filepath):
    """Parse a Word document and extract text."""
    try:
        doc = Document(filepath)
        text = []
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text.strip())
        
        # Also get tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text.append(" | ".join(row_text))
        
        return "\n".join(text)
    except Exception as e:
        print(f"Error parsing {filepath}: {e}")
        return None


def main():
    print("=" * 70)
    print("PARSING DR. HUNTER'S METHODOLOGY DOCUMENTS")
    print("=" * 70)
    
    # Find all Reference folders
    reference_folders = []
    for item in os.listdir(CME_DIR):
        if item.startswith("Reference") and "Oregon Hunter" in item:
            reference_folders.append(item)
    
    reference_folders.sort()
    print(f"\nFound {len(reference_folders)} Reference folders")
    
    knowledge_base = {
        "source": "Dr. Hunter Methodology Documents",
        "description": "Comprehensive CME examination methodology",
        "documents": []
    }
    
    for folder in reference_folders:
        folder_path = os.path.join(CME_DIR, folder)
        
        # Extract the letter and topic from folder name
        # e.g., "Reference D - Motor exam Oregon Hunter" -> D, Motor exam
        parts = folder.replace(" Oregon Hunter", "").split(" - ")
        if len(parts) >= 2:
            letter = parts[0].replace("Reference ", "")
            topic = parts[1]
        else:
            letter = "?"
            topic = folder
        
        print(f"\n[{letter}] {topic}")
        
        # Find docx files in the folder
        docx_files = []
        for f in os.listdir(folder_path):
            if f.endswith('.docx') and not f.startswith('~'):
                docx_files.append(os.path.join(folder_path, f))
        
        for docx_path in docx_files:
            content = parse_docx(docx_path)
            if content:
                doc_entry = {
                    "reference_letter": letter,
                    "topic": topic,
                    "filename": os.path.basename(docx_path),
                    "path": docx_path,
                    "content_length": len(content),
                    "content": content[:50000]  # Limit to 50k chars per doc
                }
                knowledge_base["documents"].append(doc_entry)
                print(f"   ✓ {os.path.basename(docx_path)} ({len(content)} chars)")
    
    # Save knowledge base
    output_path = "/Users/hammadhaque/Documents/cme-analysis-platform/hunter_methodology_knowledge_base.json"
    with open(output_path, 'w') as f:
        json.dump(knowledge_base, f, indent=2)
    
    print(f"\n{'='*70}")
    print(f"SUMMARY")
    print(f"{'='*70}")
    print(f"Documents parsed: {len(knowledge_base['documents'])}")
    print(f"Total content: {sum(d['content_length'] for d in knowledge_base['documents'])} chars")
    print(f"Saved to: {output_path}")
    
    # Also create a summary of key requirements per exam type
    print(f"\n{'='*70}")
    print("KEY METHODOLOGY SUMMARIES")
    print(f"{'='*70}")
    
    summaries = {}
    for doc in knowledge_base["documents"]:
        topic = doc["topic"]
        content = doc["content"][:5000]  # First 5k chars for summary
        
        # Extract key numbers and requirements
        if "24" in content or "muscle group" in content.lower():
            print(f"\n{doc['reference_letter']}. {topic}:")
            # Find lines with numbers
            for line in content.split('\n')[:30]:
                if any(num in line for num in ['24', '14', '48', '38', 'muscle', 'dermatome', 'cranial']):
                    print(f"   {line[:100]}")


if __name__ == "__main__":
    main()
